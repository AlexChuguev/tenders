from __future__ import annotations

from datetime import datetime
import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from tender_agent.analysis import TenderAnalyzer
from tender_agent.analysis_types import AnalysisPayload
from tender_agent.artifact_reader import iter_artifacts, unique_latest_artifacts
from tender_agent.config import Settings
from tender_agent.document_facts import extract_document_facts
from tender_agent.evidence import build_fact_columns, render_summary_points_with_evidence
from tender_agent.triage_rules import postprocess_payload
from tender_agent.versioning import ARTIFACT_SCHEMA_VERSION, build_versions_payload


def main() -> int:
    if len(sys.argv) != 4:
        print(
            "Usage: .venv/bin/python rebuild_results_from_current_rules.py "
            "<prepared_batch_dir> <source_xls> <output_xlsx>"
        )
        return 2

    prepared_batch_dir = Path(sys.argv[1]).expanduser().resolve()
    source_xls = Path(sys.argv[2]).expanduser().resolve()
    output_xlsx = Path(sys.argv[3]).expanduser().resolve()

    base = Path("/Users/alexchuguev/Documents/tenders")
    settings = Settings.load(base)

    analyzer = TenderAnalyzer(
        provider_name=settings.llm_provider,
        api_key=settings.llm_api_key,
        model=settings.llm_model,
        base_url=settings.llm_base_url,
        prompt_template_path=settings.prompt_template_path,
        max_chars_per_file=settings.analysis_max_chars_per_file,
        triage_policy_path=settings.triage_policy_path,
    )

    artifacts = unique_latest_artifacts(
        iter_artifacts(
            settings.artifact_dir,
            source="local_review",
            batch_name=prepared_batch_dir.name,
        )
    )
    workbook = load_workbook(output_xlsx)
    changed = 0
    skipped = 0

    folder_map = _load_folder_map(prepared_batch_dir)
    for artifact in artifacts:
        if artifact.result.error in {"deadline_passed", "no_files", "network_error"}:
            skipped += 1
            continue
        file_paths = []
        folder = _resolve_folder(prepared_batch_dir, folder_map, artifact.result.tender_id)
        if folder:
            file_paths = _collect_folder_files(folder)
        if not file_paths:
            file_paths = [Path(item) for item in artifact.result.downloaded_files if Path(item).exists()]
        if not file_paths:
            skipped += 1
            continue

        prepared_files = analyzer._prepare_files(file_paths)
        facts = extract_document_facts(
            prepared_files,
            tender_title=artifact.result.title,
            tender_price_rub=artifact.result.nmck_rub,
            policy=analyzer.triage_policy,
        )
        analysis_payload = artifact.analysis or {}
        payload = AnalysisPayload(
            decision=str(analysis_payload.get("decision") or artifact.result.decision or "Уточнить"),
            confidence_percent=int(analysis_payload.get("confidence_percent") or artifact.result.confidence_percent or 0),
            summary_points=list(analysis_payload.get("summary_points") or []),
            analysis_markdown=str(analysis_payload.get("analysis_markdown") or artifact.result.analysis_markdown or ""),
            completeness_label=str(analysis_payload.get("completeness_label") or artifact.result.document_completeness or facts.completeness_label),
            llm_raw_text=str(analysis_payload.get("llm_raw_text") or ""),
            error_type=str(analysis_payload.get("error_type") or ""),
        )
        new_payload = postprocess_payload(payload, facts, policy=analyzer.triage_policy)
        deadline_at = _parse_deadline_at(artifact.result.deadline_at)
        new_payload.summary_points = render_summary_points_with_evidence(
            facts=facts,
            decision=new_payload.decision,
            files=file_paths,
            deadline_at=deadline_at,
            summary_points=new_payload.summary_points,
        )

        artifact_json = json.loads(artifact.path.read_text(encoding="utf-8"))
        artifact_json["schema_version"] = ARTIFACT_SCHEMA_VERSION
        artifact_json["policy"] = analyzer.policy_metadata
        artifact_json["versions"] = build_versions_payload(policy_sha1=str(analyzer.policy_metadata.get("sha1") or ""))
        artifact_json["analysis"]["decision"] = new_payload.decision
        artifact_json["analysis"]["confidence_percent"] = new_payload.confidence_percent
        artifact_json["analysis"]["summary_points"] = new_payload.summary_points
        artifact_json["analysis"]["completeness_label"] = new_payload.completeness_label
        artifact_json["analysis"]["facts"] = {
            "procurement_type": facts.procurement_type,
            "key_files": facts.key_files,
            "stack": facts.stack,
            "turnover_requirements": facts.turnover_requirements,
            "project_requirements": facts.project_requirements,
            "team_requirements": facts.team_requirements,
            "licenses": facts.licenses,
            "payment": facts.payment,
            "document_roles": facts.document_roles,
            "completeness_label": facts.completeness_label,
            "completeness_notes": facts.completeness_notes,
            "triage_signals": facts.triage_signals,
            "participant_requirements": {
                "requires_sro": facts.participant_requirements.requires_sro,
                "requires_analog_projects": facts.participant_requirements.requires_analog_projects,
                "requires_existing_system_modification_experience": facts.participant_requirements.requires_existing_system_modification_experience,
                "requires_financial_documents": facts.participant_requirements.requires_financial_documents,
                "requires_team_documents": facts.participant_requirements.requires_team_documents,
                "min_turnover_rub": facts.participant_requirements.min_turnover_rub,
                "turnover_texts": facts.participant_requirements.turnover_texts,
                "experience_texts": facts.participant_requirements.experience_texts,
                "required_roles": facts.participant_requirements.required_roles,
                "licenses_or_certs": facts.participant_requirements.licenses_or_certs,
                "other_requirements": facts.participant_requirements.other_requirements,
                "submission_docs": facts.participant_requirements.submission_docs,
            },
            "submission_requirements": facts.submission_requirements,
            "quality_flags": facts.quality_flags,
            "reason_codes": facts.reason_codes,
        }
        artifact_json["result"]["decision"] = new_payload.decision
        artifact_json["result"]["confidence_percent"] = new_payload.confidence_percent
        artifact_json["result"]["summary_text"] = "\n".join(new_payload.summary_points)
        artifact_json["result"]["document_completeness"] = new_payload.completeness_label
        facts_stack, facts_requirements, facts_docs, facts_payment = build_fact_columns(facts)
        artifact_json["result"]["facts_stack"] = facts_stack
        artifact_json["result"]["facts_requirements"] = facts_requirements
        artifact_json["result"]["facts_docs"] = facts_docs
        artifact_json["result"]["facts_payment"] = facts_payment
        artifact_json["result"]["reason_codes"] = list(facts.reason_codes)
        artifact.path.write_text(json.dumps(artifact_json, ensure_ascii=False, indent=2), encoding="utf-8")

        if _update_excel_row(
            workbook,
            url=artifact.result.url,
            title=artifact.result.title,
            decision=new_payload.decision,
            confidence=new_payload.confidence_percent,
            summary_text="\n".join(new_payload.summary_points),
        ):
            changed += 1
        else:
            skipped += 1

    workbook.save(output_xlsx)
    print(f"changed={changed}")
    print(f"skipped={skipped}")
    return 0


def _load_folder_map(prepared_batch_dir: Path) -> dict[str, str]:
    manifest = prepared_batch_dir / "_manifest.csv"
    if not manifest.exists():
        return {}
    mapping: dict[str, str] = {}
    with manifest.open(encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            tender_id = str(row.get("tender_id") or "").strip()
            folder_name = str(row.get("folder_name") or "").strip()
            if tender_id and folder_name:
                mapping[tender_id] = folder_name
    return mapping


def _resolve_folder(prepared_batch_dir: Path, folder_map: dict[str, str], tender_id: str) -> Path | None:
    if not tender_id:
        return None
    direct = prepared_batch_dir / tender_id
    if direct.exists():
        return direct
    folder_name = folder_map.get(tender_id)
    if folder_name:
        candidate = prepared_batch_dir / folder_name
        if candidate.exists():
            return candidate
    return None


def _collect_folder_files(folder: Path) -> list[Path]:
    files: list[Path] = []
    for path in folder.rglob("*"):
        if not path.is_file():
            continue
        if path.name.startswith(".") or path.name == ".done":
            continue
        if "__extracted" not in str(path) and path.suffix.lower() in {".zip", ".rar", ".7z"}:
            continue
        files.append(path)
    return sorted(files)


def _update_excel_row(workbook, *, url: str, title: str, decision: str, confidence: int, summary_text: str) -> bool:
    target_url = (url or "").strip()
    target_title = (title or "").strip()
    for worksheet in workbook.worksheets:
        for row_index in range(2, worksheet.max_row + 1):
            title_cell = worksheet.cell(row=row_index, column=4)
            row_title = str(title_cell.value or "").strip()
            row_url = title_cell.hyperlink.target if title_cell.hyperlink else ""
            if target_url and row_url == target_url:
                worksheet.cell(row=row_index, column=3).value = summary_text
                return True
            if target_title and row_title == target_title:
                worksheet.cell(row=row_index, column=3).value = summary_text
                return True
    return False


def _parse_deadline_at(value: str | None) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value))
    except Exception:
        return None


if __name__ == "__main__":
    raise SystemExit(main())
