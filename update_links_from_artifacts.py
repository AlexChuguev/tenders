from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from tender_agent.artifact_reader import iter_artifacts
from tender_agent.config import Settings
from tender_agent.export.excel_schema import column_index


def main() -> int:
    if len(sys.argv) < 2:
        print(
            "Usage: .venv/bin/python update_links_from_artifacts.py /path/to/tender_analysis.xlsx [batch_name]"
        )
        return 2

    workbook_path = Path(sys.argv[1]).expanduser().resolve()
    batch_name = sys.argv[2].strip() if len(sys.argv) > 2 else ""
    if not workbook_path.exists():
        print(f"Excel not found: {workbook_path}")
        return 2

    base = Path("/Users/alexchuguev/Documents/tenders")
    settings = Settings.load(base)

    artifacts = iter_artifacts(
        settings.artifact_dir,
        batch_name=batch_name or None,
    )
    by_url = {a.result.url.strip(): a for a in artifacts if a.result.url}
    by_id = {a.result.tender_id.strip(): a for a in artifacts if a.result.tender_id}

    wb = load_workbook(workbook_path)
    updated = 0

    for ws in wb.worksheets:
        for row in range(2, ws.max_row + 1):
            title_cell = ws.cell(row=row, column=column_index("title"))
            url = title_cell.hyperlink.target.strip() if title_cell.hyperlink and title_cell.hyperlink.target else ""
            artifact = by_url.get(url) if url else None
            if artifact is None and url and "/tender/" in url:
                tender_id = url.rsplit("/", 1)[-1].strip()
                if tender_id:
                    artifact = by_id.get(tender_id)
            if artifact is None:
                continue

            summary_cell = ws.cell(row=row, column=column_index("summary"))
            if not summary_cell.hyperlink and artifact.result.downloaded_files:
                file_path = Path(artifact.result.downloaded_files[0]).expanduser().resolve()
                if file_path.exists():
                    summary_cell.hyperlink = file_path.as_uri()
                    summary_cell.style = "Hyperlink"
                    updated += 1

            batch_cell = ws.cell(row=row, column=column_index("batch"))
            batch_path = artifact.result.batch_path
            if not batch_path and artifact.result.batch_name:
                candidate = base / "manual_downloads" / artifact.result.batch_name
                if candidate.exists():
                    batch_path = str(candidate)
            if batch_path and not batch_cell.hyperlink and batch_cell.value:
                batch_dir = Path(batch_path).expanduser().resolve()
                if batch_dir.exists():
                    batch_cell.hyperlink = batch_dir.as_uri()
                    batch_cell.style = "Hyperlink"
                    updated += 1

    wb.save(workbook_path)
    print(f"updated_links={updated}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
