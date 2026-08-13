from __future__ import annotations

import csv
import sys
from collections import Counter
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit(
            "Usage: .venv/bin/python seldon_external_host_stats.py /path/to/history.xlsx [limit] [output_csv]"
        )

    workbook_path = Path(sys.argv[1]).expanduser().resolve()
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 100
    output_csv = (
        Path(sys.argv[3]).expanduser().resolve()
        if len(sys.argv) > 3
        else workbook_path.with_name("seldon_external_hosts_sample.csv")
    )

    if not workbook_path.exists():
        raise SystemExit(f"Workbook not found: {workbook_path}")

    urls = _extract_seldon_urls(workbook_path, limit)
    if not urls:
        raise SystemExit("No Seldon tender URLs found in workbook.")

    profile_dir = Path(__file__).resolve().parent / "chrome_profiles" / "seldon"
    profile_dir.mkdir(parents=True, exist_ok=True)

    rows: list[dict[str, str]] = []
    host_counts: Counter[str] = Counter()

    with sync_playwright() as playwright:
        context = playwright.chromium.launch_persistent_context(
            user_data_dir=str(profile_dir),
            channel="chrome",
            headless=False,
        )
        context.set_default_timeout(12000)
        context.set_default_navigation_timeout(18000)
        try:
            page = context.new_page()
            for index, url in enumerate(urls, start=1):
                result = _inspect_tender(page, url)
                rows.append(result)
                host = result["external_host"]
                if host:
                    host_counts[host] += 1
                print(
                    f"[{index}/{len(urls)}] tender={result['tender_id']} external_host={host or '-'} status={result['status']}",
                    flush=True,
                )
        finally:
            context.close()

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=["tender_url", "tender_id", "external_url", "external_host", "status"],
        )
        writer.writeheader()
        writer.writerows(rows)

    print("TOP_HOSTS", flush=True)
    for host, count in host_counts.most_common():
        print(f"{host}\t{count}", flush=True)
    print(f"CSV: {output_csv}", flush=True)


def _extract_seldon_urls(workbook_path: Path, limit: int) -> list[str]:
    wb = load_workbook(workbook_path, read_only=False, data_only=False)
    seen: set[str] = set()
    ordered: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                hyperlink = cell.hyperlink.target if cell.hyperlink else None
                if not hyperlink or "pro.myseldon.com/ru/tender/" not in hyperlink:
                    continue
                if hyperlink in seen:
                    continue
                seen.add(hyperlink)
                ordered.append(hyperlink)
                if len(ordered) >= limit:
                    return ordered
    return ordered


def _inspect_tender(page, tender_url: str) -> dict[str, str]:
    try:
        page.goto(tender_url, wait_until="domcontentloaded", timeout=18000)
        try:
            page.wait_for_load_state("networkidle", timeout=7000)
        except Exception:
            pass
    except Exception as exc:
        return {
            "tender_url": tender_url,
            "tender_id": _tender_id_from_url(tender_url),
            "external_url": "",
            "external_host": "",
            "status": f"goto_error: {type(exc).__name__}",
        }

    links = []
    for locator in page.locator("a").all():
        href = locator.get_attribute("href")
        if not href:
            continue
        absolute = page.evaluate(
            """(href) => {
                try { return new URL(href, window.location.href).toString(); }
                catch { return href; }
            }""",
            href,
        )
        host = urlparse(absolute).netloc.lower()
        if not host or "myseldon.com" in host or "basis.myseldon.com" in host:
            continue
        if any(token in host for token in ("telegram.me", "t.me", "vk.cc", "youtube.com", "youtu.be")):
            continue
        links.append(absolute)

    external_url = _pick_external_link(links)
    return {
        "tender_url": tender_url,
        "tender_id": _tender_id_from_url(tender_url),
        "external_url": external_url,
        "external_host": urlparse(external_url).netloc.lower(),
        "status": "ok" if external_url else "no_external_link",
    }


def _pick_external_link(links: list[str]) -> str:
    preferred_tokens = ("tender", "zakup", "purchase", "process", "trade", "view", "notice", "auction")
    for link in links:
        lower = link.lower()
        if any(token in lower for token in preferred_tokens):
            return link
    return links[0] if links else ""


def _tender_id_from_url(url: str) -> str:
    parts = [part for part in urlparse(url).path.split("/") if part]
    return parts[-1] if parts else ""


if __name__ == "__main__":
    main()
