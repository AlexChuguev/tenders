from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from tender_agent.artifact_reader import iter_artifacts
from tender_agent.excel_loader import load_tenders
from tender_agent.export.excel_schema import column_index


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit(
            "Usage: .venv/bin/python repair_excel_hyperlinks.py /path/to/output.xlsx /path/to/export.xls /path/to/artifacts <batch_name>"
        )

    workbook_path = Path(sys.argv[1]).expanduser().resolve()
    xls_path = Path(sys.argv[2]).expanduser().resolve()
    artifact_root = Path(sys.argv[3]).expanduser().resolve()
    batch_name = sys.argv[4].strip()

    title_to_url: dict[str, str] = {}

    for artifact in iter_artifacts(artifact_root, source="local_review", batch_name=batch_name):
        title = artifact.result.title.strip()
        url = artifact.result.url.strip()
        if title and url:
            title_to_url[title] = url

    for tender in load_tenders(
        xls_path=xls_path,
        url_column="Номер извещения (ссылка на источник)",
        id_column="Номер извещения (ссылка на источник)",
        title_column="Наименование лота",
    ):
        title = tender.title.strip()
        url = tender.url.strip()
        if title and url and title not in title_to_url:
            title_to_url[title] = url

    workbook = load_workbook(workbook_path)
    restored = 0
    unresolved: list[tuple[str, int, str]] = []
    for worksheet in workbook.worksheets:
        for row_index in range(5, worksheet.max_row + 1):
            title_cell = worksheet.cell(row=row_index, column=column_index("title"))
            title = str(title_cell.value or "").strip()
            if not title:
                continue
            hyperlink = title_cell.hyperlink
            if hyperlink and hyperlink.target:
                continue
            url = title_to_url.get(title, "").strip()
            if not url:
                unresolved.append((worksheet.title, row_index, title))
                continue
            title_cell.hyperlink = url
            title_cell.style = "Hyperlink"
            restored += 1

    workbook.save(workbook_path)
    print(f"restored={restored}")
    print(f"unresolved={len(unresolved)}")
    for sheet_name, row_index, title in unresolved[:20]:
        print(f"{sheet_name}:{row_index}: {title}")


if __name__ == "__main__":
    main()
