from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tender_agent.export.excel_migration import (
    cleanup_legacy_workbook,
    normalize_deadline_cells,
    remove_malformed_rows,
    remove_technical_rows,
    renumber_rows,
    sort_sheet_by_deadline,
    is_technical_retry_result,
)
from tender_agent.export.excel_schema import (
    apply_deadline_format,
    apply_status_fill,
    column_index,
    ensure_sheet_structure,
    set_widths,
    sheet_name,
)
from tender_agent.models import TenderAnalysisResult


class ExcelWriter:
    def __init__(self, output_path: Path, *, preserve_status: bool = False) -> None:
        self.output_path = output_path
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        self.preserve_status = preserve_status
        self._backup_created = False

    def ensure_header(self) -> None:
        workbook = self._open()
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.calcMode = "auto"
        cleanup_legacy_workbook(workbook)
        remove_technical_rows(workbook)
        remove_malformed_rows(workbook)
        normalize_deadline_cells(workbook)
        for worksheet in workbook.worksheets:
            ensure_sheet_structure(worksheet)
            _clear_status_fills(worksheet)
            renumber_rows(worksheet)
            set_widths(worksheet)
            sort_sheet_by_deadline(worksheet)
        self._save(workbook)

    def append_result(self, result: TenderAnalysisResult) -> None:
        if is_technical_retry_result(result.summary_text, result.confidence_percent) and str(
            result.error or ""
        ).strip() != "network_error":
            return
        workbook = self._open()
        cleanup_legacy_workbook(workbook)
        remove_technical_rows(workbook)
        remove_malformed_rows(workbook)
        normalize_deadline_cells(workbook)
        worksheet = self._ensure_month_sheet(workbook, result.deadline_at)
        target_row = _find_existing_row(workbook, result)
        status_value = _derive_status(result)
        if self.preserve_status:
            status_value = ""
        if target_row is not None:
            target_sheet, row_index = target_row
            existing_status = str(
                target_sheet.cell(row=row_index, column=column_index("status")).value or ""
            ).strip()
            status_to_set = _merge_status(existing_status, status_value)
            summary_cell = target_sheet.cell(
                row=row_index, column=column_index("summary"), value=result.summary_text
            )
            target_sheet.cell(row=row_index, column=column_index("title"), value=result.title)
            target_sheet.cell(row=row_index, column=column_index("deadline"), value=result.deadline_at)
            batch_cell = target_sheet.cell(row=row_index, column=column_index("batch"))
            if result.batch_name or batch_cell.value:
                batch_cell.value = result.batch_name or batch_cell.value
            _apply_summary_hyperlink(summary_cell, result.downloaded_files)
            _apply_batch_hyperlink(batch_cell, result.batch_path)
            _write_facts_cells(target_sheet, row_index, result)
            if status_to_set:
                target_sheet.cell(row=row_index, column=column_index("status"), value=status_to_set)
                apply_status_fill(
                    target_sheet.cell(row=row_index, column=column_index("status"))
                )
            title_cell = target_sheet.cell(row=row_index, column=column_index("title"))
            title_cell.hyperlink = result.url
            title_cell.style = "Hyperlink"
            apply_deadline_format(
                target_sheet.cell(row=row_index, column=column_index("deadline"))
            )
        else:
            row_number = max(worksheet.max_row + 1, 2)
            worksheet.append(
                [
                    "",
                    status_value or "",
                    result.summary_text,
                    result.title,
                    result.deadline_at,
                    "",
                    result.batch_name or "",
                    result.facts_stack or "",
                    result.facts_requirements or "",
                    result.facts_docs or "",
                    result.facts_payment or "",
                    result.facts_req_sro or "",
                    result.facts_req_turnover or "",
                    result.facts_req_analog_projects or "",
                    result.facts_req_legacy_experience or "",
                    result.facts_req_roles or "",
                    result.facts_req_licenses or "",
                ]
            )
            title_cell = worksheet.cell(row=worksheet.max_row, column=column_index("title"))
            title_cell.hyperlink = result.url
            title_cell.style = "Hyperlink"
            summary_cell = worksheet.cell(row=worksheet.max_row, column=column_index("summary"))
            _apply_summary_hyperlink(summary_cell, result.downloaded_files)
            batch_cell = worksheet.cell(row=worksheet.max_row, column=column_index("batch"))
            _apply_batch_hyperlink(batch_cell, result.batch_path)
            apply_deadline_format(
                worksheet.cell(row=worksheet.max_row, column=column_index("deadline"))
            )
            if status_value:
                apply_status_fill(
                    worksheet.cell(row=worksheet.max_row, column=column_index("status"))
                )
        for sheet in workbook.worksheets:
            renumber_rows(sheet)
            sort_sheet_by_deadline(sheet)
        self._restore_title_hyperlink(worksheet, result.title, result.url, result.deadline_at)
        self._save(workbook)

    def _open(self):
        if self.output_path.exists():
            return load_workbook(self.output_path)
        workbook = Workbook()
        workbook.active.title = "Temp"
        return workbook

    def _save(self, workbook) -> None:
        self._backup_once()
        workbook.save(self.output_path)

    def _backup_once(self) -> None:
        if self._backup_created:
            return
        self._backup_created = True
        if not self.output_path.exists():
            return
        backup_path = self.output_path.with_name(f"{self.output_path.stem}.backup{self.output_path.suffix}")
        shutil.copy2(self.output_path, backup_path)

    def _ensure_month_sheet(self, workbook, deadline_at: datetime | None):
        target_sheet_name = sheet_name(deadline_at)
        if target_sheet_name in workbook.sheetnames:
            worksheet = workbook[target_sheet_name]
            ensure_sheet_structure(worksheet)
            return worksheet

        worksheet = workbook.create_sheet(title=target_sheet_name)
        if "Temp" in workbook.sheetnames and len(workbook.sheetnames) > 1:
            del workbook["Temp"]
        ensure_sheet_structure(worksheet)
        return worksheet

    @staticmethod
    def _restore_title_hyperlink(worksheet, title: str, url: str, deadline_at: datetime | None) -> None:
        if not title or not url:
            return
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=4)
            if str(cell.value or "").strip() != title:
                continue
            row_deadline = worksheet.cell(row=row_index, column=5).value
            if deadline_at and row_deadline and row_deadline != deadline_at:
                continue
            if cell.hyperlink and cell.hyperlink.target:
                continue
            cell.hyperlink = url
            cell.style = "Hyperlink"


def _find_existing_row(workbook, result: TenderAnalysisResult):
    target_url = (result.url or "").strip()
    target_title = (result.title or "").strip()
    tender_id = (result.tender_id or "").strip()
    batch_name = (result.batch_name or "").strip()
    for worksheet in workbook.worksheets:
        for row_index in range(2, worksheet.max_row + 1):
            title_cell = worksheet.cell(row=row_index, column=column_index("title"))
            row_title = str(title_cell.value or "").strip()
            row_url = title_cell.hyperlink.target if title_cell.hyperlink else ""
            row_batch = str(
                worksheet.cell(row=row_index, column=column_index("batch")).value or ""
            ).strip()
            if target_url and row_url == target_url:
                return worksheet, row_index
            if tender_id and row_url and tender_id in row_url:
                return worksheet, row_index
            if target_title and row_title == target_title and not target_url and not tender_id:
                return worksheet, row_index
    return None


def _derive_status(result: TenderAnalysisResult) -> str:
    error = str(result.error or "").strip()
    summary = str(result.summary_text or "").strip()
    if error == "no_files" or summary.startswith("Файлы не добавлены"):
        return "Документы не загружены"
    if error in {"network_error", "llm_error"} or summary.startswith("Техсбой LLM"):
        return "Техсбой LLM"
    if summary:
        return "Системой обработано"
    return ""


def _merge_status(existing: str, computed: str) -> str:
    if not computed:
        return existing
    if not existing:
        return computed
    if existing in {"Документы не загружены", "Техсбой LLM", "LLM", "Системой обработано"}:
        return computed
    return existing


def _clear_status_fills(worksheet) -> None:
    from openpyxl.styles import PatternFill

    no_fill = PatternFill(fill_type=None)
    for row_index in range(2, worksheet.max_row + 1):
        worksheet.cell(row=row_index, column=column_index("status")).fill = no_fill


def _write_facts_cells(worksheet, row_index: int, result: TenderAnalysisResult) -> None:
    worksheet.cell(row=row_index, column=column_index("facts_stack"), value=result.facts_stack or "")
    worksheet.cell(
        row=row_index, column=column_index("facts_requirements"), value=result.facts_requirements or ""
    )
    worksheet.cell(row=row_index, column=column_index("facts_docs"), value=result.facts_docs or "")
    worksheet.cell(row=row_index, column=column_index("facts_payment"), value=result.facts_payment or "")
    worksheet.cell(row=row_index, column=column_index("facts_req_sro"), value=result.facts_req_sro or "")
    worksheet.cell(row=row_index, column=column_index("facts_req_turnover"), value=result.facts_req_turnover or "")
    worksheet.cell(
        row=row_index,
        column=column_index("facts_req_analog_projects"),
        value=result.facts_req_analog_projects or "",
    )
    worksheet.cell(
        row=row_index,
        column=column_index("facts_req_legacy_experience"),
        value=result.facts_req_legacy_experience or "",
    )
    worksheet.cell(row=row_index, column=column_index("facts_req_roles"), value=result.facts_req_roles or "")
    worksheet.cell(
        row=row_index,
        column=column_index("facts_req_licenses"),
        value=result.facts_req_licenses or "",
    )


def _apply_summary_hyperlink(cell, downloaded_files: list[str]) -> None:
    if cell is None or cell.value in (None, ""):
        return
    if cell.hyperlink and cell.hyperlink.target:
        return
    if not downloaded_files:
        return
    try:
        path = Path(downloaded_files[0]).expanduser().resolve()
    except Exception:
        return
    if not path.exists():
        return
    cell.hyperlink = path.as_uri()
    cell.style = "Hyperlink"


def _apply_batch_hyperlink(cell, batch_path: str) -> None:
    if cell is None or cell.value in (None, ""):
        return
    if cell.hyperlink and cell.hyperlink.target:
        return
    if not batch_path:
        return
    try:
        path = Path(batch_path).expanduser().resolve()
    except Exception:
        return
    if not path.exists():
        return
    cell.hyperlink = path.as_uri()
    cell.style = "Hyperlink"
