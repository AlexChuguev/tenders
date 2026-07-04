from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from tender_agent.models import TenderAnalysisResult


@dataclass(frozen=True)
class ColumnSpec:
    key: str
    header: str
    width_px: int
    owner: str
    editable: bool = False


COLUMNS = [
    ColumnSpec("number", "№", 30, owner="system"),
    ColumnSpec("status", "Статус", 140, owner="user", editable=True),
    ColumnSpec("summary", "Summary", 369, owner="system"),
    ColumnSpec("title", "Наименование лота", 369, owner="system"),
    ColumnSpec("deadline", "Дата окончания приема заявок", 125, owner="system"),
    ColumnSpec("comment", "Комментарий", 240, owner="user", editable=True),
    ColumnSpec("batch", "Выгрузка", 200, owner="system"),
    ColumnSpec("facts_stack", "Факты: стек", 220, owner="system"),
    ColumnSpec("facts_requirements", "Факты: требования", 260, owner="system"),
    ColumnSpec("facts_docs", "Факты: документация", 220, owner="system"),
    ColumnSpec("facts_payment", "Факты: оплата", 220, owner="system"),
    ColumnSpec("facts_req_sro", "Факт: СРО", 90, owner="system"),
    ColumnSpec("facts_req_turnover", "Факт: оборот", 140, owner="system"),
    ColumnSpec("facts_req_analog_projects", "Факт: аналогичный опыт", 130, owner="system"),
    ColumnSpec("facts_req_legacy_experience", "Факт: опыт доработки", 150, owner="system"),
    ColumnSpec("facts_req_roles", "Факт: роли", 180, owner="system"),
    ColumnSpec("facts_req_licenses", "Факт: лицензии/сертификаты", 180, owner="system"),
]

DETAIL_HEADERS = [column.header for column in COLUMNS]
HEADER_TO_KEY = {column.header: column.key for column in COLUMNS}
KEY_TO_INDEX = {column.key: index + 1 for index, column in enumerate(COLUMNS)}

MONTH_NAMES = {
    1: "Январь",
    2: "Февраль",
    3: "Март",
    4: "Апрель",
    5: "Май",
    6: "Июнь",
    7: "Июль",
    8: "Август",
    9: "Сентябрь",
    10: "Октябрь",
    11: "Ноябрь",
    12: "Декабрь",
}

STATUS_FONT_COLORS = {
    "Документы не загружены": "808080",
    "Техсбой LLM": "B8860B",
    "Системой обработано": "1F4E79",
    "LLM": "B8860B",
    "Проверено": "008000",
    "В процессе": "7030A0",
    "подано": "000000",
    "Подано": "000000",
    "Не получилось скачать документы": "800000",
    "в минус": "0070C0",
    "Иное": "00B0F0",
}


def sheet_name(deadline_at: datetime | None) -> str:
    if deadline_at is None:
        return "Без даты"
    return f"{MONTH_NAMES[deadline_at.month]} {deadline_at.year}"


def ensure_sheet_structure(worksheet) -> None:
    for index, header in enumerate(DETAIL_HEADERS, start=1):
        worksheet.cell(row=1, column=index, value=header)
    _ensure_status_validation(worksheet)
    _ensure_status_conditional_formatting(worksheet)
    set_widths(worksheet)


def set_widths(worksheet) -> None:
    for index, column in enumerate(COLUMNS, start=1):
        column_letter = worksheet.cell(row=1, column=index).column_letter
        worksheet.column_dimensions[column_letter].width = excel_width_from_pixels(column.width_px)

def apply_status_fill(cell) -> None:
    cell.fill = PatternFill(fill_type=None)
    cell.font = Font(color=None)


def _ensure_status_validation(worksheet) -> None:
    if getattr(worksheet, "data_validations", None) is not None:
        to_keep = []
        for dv in worksheet.data_validations.dataValidation:
            if dv.type == "list" and "B" in str(dv.sqref):
                continue
            to_keep.append(dv)
        worksheet.data_validations.dataValidation = to_keep
    options = [
        "Документы не загружены",
        "Техсбой LLM",
        "Системой обработано",
        "LLM",
        "Проверено",
        "В процессе",
        "Подано",
        "Не получилось скачать документы",
        "в минус",
        "Иное",
    ]
    dv = DataValidation(type="list", formula1=f"\"{','.join(options)}\"", allow_blank=True)
    worksheet.add_data_validation(dv)
    dv.add("B2:B1048576")


def _ensure_status_conditional_formatting(worksheet) -> None:
    try:
        rules = worksheet.conditional_formatting._cf_rules
        for key in list(rules.keys()):
            if "B2" in str(key):
                del rules[key]
    except Exception:
        pass
    for label, color in STATUS_FONT_COLORS.items():
        rule = CellIsRule(
            operator="equal",
            formula=[f'"{label}"'],
            font=Font(color=color),
            stopIfTrue=False,
        )
        worksheet.conditional_formatting.add("B2:B1048576", rule)


def apply_deadline_format(cell) -> None:
    value = cell.value
    if not isinstance(value, datetime):
        return
    if value.hour == 0 and value.minute == 0 and value.second == 0 and value.microsecond == 0:
        cell.number_format = "D.M.YYYY"
        return
    cell.number_format = "D.M.YYYY HH:MM"


def normalize_decision_for_sheet(result: TenderAnalysisResult) -> str:
    normalized = str(result.decision or "").strip()
    if normalized:
        return normalized
    return "Уточнить"


def coerce_deadline_value(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        normalized = value.strip()
        if not normalized:
            return None
        for fmt in ("%d.%m.%Y %H:%M", "%d.%m.%Y"):
            try:
                return datetime.strptime(normalized, fmt)
            except ValueError:
                continue
    return None


def excel_width_from_pixels(pixels: int) -> float:
    if pixels <= 12:
        return float(pixels)
    return round((pixels - 5) / 7, 2)


def column_index(key: str) -> int:
    return KEY_TO_INDEX[key]


def column_count() -> int:
    return len(COLUMNS)


def system_owned_keys() -> tuple[str, ...]:
    return tuple(column.key for column in COLUMNS if column.owner == "system")


def user_owned_keys() -> tuple[str, ...]:
    return tuple(column.key for column in COLUMNS if column.owner == "user")


def map_legacy_tag_to_decision(tag: str) -> str:
    normalized = tag.strip()
    if normalized in {"Предварительно подходит", "В работу", "Подано"}:
        return "Брать"
    if normalized == "Не изучено":
        return "Уточнить"
    return "Не брать"
