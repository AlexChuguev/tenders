from __future__ import annotations

import csv
import json
from dataclasses import asdict, dataclass
from pathlib import Path

from tender_agent.artifact_reader import iter_artifacts
from tender_agent.config import Settings
from tender_agent.excel_loader import load_tenders


@dataclass(frozen=True)
class BatchRowAudit:
    row_number: int
    tender_id: str
    title: str
    url: str
    folder: str
    files: int
    artifact: bool
    issue: str


@dataclass(frozen=True)
class BatchAuditReport:
    batch_name: str
    xls_rows: int
    folders_found: int
    folders_with_files: int
    artifacts_written: int
    excel_rows: int
    issues_total: int
    health: str
    rows: list[BatchRowAudit]


def audit_batch(
    *,
    batch_dir: Path,
    xls_path: Path,
    artifacts_root: Path,
    excel_path: Path | None = None,
) -> BatchAuditReport:
    base = Path("/Users/alexchuguev/Documents/tenders")
    settings = Settings.load(base)

    tenders = load_tenders(
        xls_path=xls_path,
        url_column=settings.platform_tender_url_column,
        id_column=settings.platform_tender_id_column,
        title_column=settings.platform_tender_title_column,
    )
    manifest_map = _load_manifest_map(batch_dir)
    artifacts = iter_artifacts(artifacts_root, source="local_review", batch_name=batch_dir.name)
    by_url = {a.result.url.strip(): a for a in artifacts if a.result.url}
    by_id = {a.result.tender_id.strip(): a for a in artifacts if a.result.tender_id}

    rows: list[BatchRowAudit] = []
    folders_found = 0
    folders_with_files = 0
    for row_number, tender in enumerate(tenders, start=1):
        folder = _resolve_folder(batch_dir, manifest_map, tender.tender_id, row_number)
        has_folder = folder is not None
        files = _collect_files(folder) if folder else []
        has_files = bool(files)
        has_artifact = False
        if tender.url and tender.url.strip() in by_url:
            has_artifact = True
        elif tender.tender_id and tender.tender_id.strip() in by_id:
            has_artifact = True

        if has_folder:
            folders_found += 1
        if has_files:
            folders_with_files += 1

        if has_files and not has_artifact:
            issue = "files_without_artifact"
        elif has_folder and not has_files:
            issue = "empty_folder"
        elif not has_folder:
            issue = "missing_folder"
        else:
            issue = ""

        rows.append(
            BatchRowAudit(
                row_number=row_number,
                tender_id=tender.tender_id,
                title=tender.title,
                url=tender.url,
                folder=str(folder) if folder else "",
                files=len(files),
                artifact=has_artifact,
                issue=issue,
            )
        )

    excel_rows = _count_excel_rows_for_batch(excel_path, batch_dir.name) if excel_path else 0
    issues_total = sum(1 for row in rows if row.issue)
    health = "ok"
    if any(row.issue == "files_without_artifact" for row in rows):
        health = "broken"
    elif issues_total:
        health = "incomplete"

    return BatchAuditReport(
        batch_name=batch_dir.name,
        xls_rows=len(tenders),
        folders_found=folders_found,
        folders_with_files=folders_with_files,
        artifacts_written=len(artifacts),
        excel_rows=excel_rows,
        issues_total=issues_total,
        health=health,
        rows=rows,
    )


def write_batch_audit(report: BatchAuditReport, *, batch_dir: Path) -> tuple[Path, Path]:
    csv_path = batch_dir / "_completeness_report.csv"
    json_path = batch_dir / "_completeness_report.json"
    with csv_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=["row_number", "tender_id", "title", "url", "folder", "files", "artifact", "issue"],
        )
        writer.writeheader()
        for row in report.rows:
            writer.writerow(asdict(row))
    json_path.write_text(
        json.dumps(
            {
                "batch_name": report.batch_name,
                "xls_rows": report.xls_rows,
                "folders_found": report.folders_found,
                "folders_with_files": report.folders_with_files,
                "artifacts_written": report.artifacts_written,
                "excel_rows": report.excel_rows,
                "issues_total": report.issues_total,
                "health": report.health,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return csv_path, json_path


def _load_manifest_map(batch_dir: Path) -> dict[str, str]:
    manifest = batch_dir / "_manifest.csv"
    if not manifest.exists():
        return {}
    mapping: dict[str, str] = {}
    with manifest.open(encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            tender_id = str(row.get("tender_id") or "").strip()
            folder_name = str(row.get("folder_name") or "").strip()
            if tender_id and folder_name:
                mapping[tender_id] = folder_name
    return mapping


def _resolve_folder(batch_dir: Path, manifest_map: dict[str, str], tender_id: str, row_number: int) -> Path | None:
    if tender_id:
        direct = batch_dir / tender_id
        if direct.exists():
            return direct
        if tender_id in manifest_map:
            candidate = batch_dir / manifest_map[tender_id]
            if candidate.exists():
                return candidate
    prefix = f"{row_number:03d}."
    for folder in batch_dir.iterdir():
        if folder.is_dir() and folder.name.startswith(prefix):
            return folder
    return None


def _collect_files(folder: Path | None) -> list[Path]:
    if folder is None:
        return []
    files: list[Path] = []
    for path in folder.rglob("*"):
        if not path.is_file() or path.name.startswith("."):
            continue
        try:
            if path.stat().st_size == 0:
                continue
        except OSError:
            continue
        files.append(path)
    return files


def _count_excel_rows_for_batch(excel_path: Path | None, batch_name: str) -> int:
    if excel_path is None or not excel_path.exists():
        return 0
    from openpyxl import load_workbook

    wb = load_workbook(excel_path)
    total = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) >= 7 and str(row[6] or "").strip() == batch_name:
                total += 1
    return total
