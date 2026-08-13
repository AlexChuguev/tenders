from __future__ import annotations

import json
import hashlib
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
try:
    import xlrd
except Exception:  # pragma: no cover - optional dependency in some environments
    xlrd = None
try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional dependency in some environments
    PdfReader = None

from tender_agent.analysis_types import AnalysisPayload, DECISIONS
from tender_agent.document_facts import (
    extract_document_facts,
    format_price_for_prompt,
    render_extracted_facts,
)
from tender_agent.extraction_report import build_extraction_report
from tender_agent.llm import create_llm_provider
from tender_agent.llm_findings import parse_llm_findings, verify_llm_findings
from tender_agent.policy import load_triage_policy, policy_metadata
from tender_agent.summary_builder import build_canonical_summary_points
from tender_agent.text_extract import extract_doc_with_textutil
from tender_agent.triage_rules import postprocess_payload


EXTRACTION_CACHE_VERSION = "2026-07-10-chunked-v2"


class TenderAnalyzer:
    def __init__(
        self,
        provider_name: str,
        api_key: str,
        model: str,
        base_url: str,
        prompt_template_path: Path,
        max_chars_per_file: int = 12000,
        triage_policy_path: Path | None = None,
    ) -> None:
        self.provider_name = provider_name
        self.max_chars_per_file = max_chars_per_file
        self.provider = create_llm_provider(
            provider_name=provider_name,
            api_key=api_key,
            model=model,
            base_url=base_url or None,
            max_chars_per_file=max_chars_per_file,
        )
        self.prompt_template = prompt_template_path.read_text(encoding="utf-8")
        self.triage_policy = load_triage_policy(triage_policy_path)
        self.policy_metadata = policy_metadata(self.triage_policy)
        self.extract_cache_dir = prompt_template_path.resolve().parent / "state" / "extraction_cache"
        self.extract_cache_dir.mkdir(parents=True, exist_ok=True)
        self._extraction_metadata: dict[str, dict[str, object]] = {}

    def analyze(self, tender_url: str, files: list[Path]) -> AnalysisPayload:
        return self.analyze_with_context(tender_url=tender_url, files=files)

    def analyze_with_context(
        self,
        tender_url: str,
        files: list[Path],
        tender_title: str = "",
        tender_price_rub: float | None = None,
    ) -> AnalysisPayload:
        if self.provider_name != "stub" and not files:
            return AnalysisPayload(
                decision="Уточнить",
                confidence_percent=0,
                summary_points=[
                    "Документы не получены, поэтому содержательная оценка тендера не выполнена.",
                    "Невозможно подтвердить стек, объём работ и требования к участнику.",
                    "Невозможно оценить пакет документов для участия.",
                    "Невозможно проверить условия оплаты и обеспечения.",
                    "Нужны сами тендерные материалы для принятия решения.",
                ],
                analysis_markdown="Документы не были скачаны, анализ невозможен.",
                completeness_label="низкая",
                error_type="missing_documents",
            )
        temp_dir: Path | None = None
        facts = None
        try:
            prepared_files = self._prepare_files(files)
            temp_dir = getattr(self, "_temp_dir", None)
            extraction_report = build_extraction_report(
                source_files=files,
                prepared_files=prepared_files,
                metadata_by_prepared=self._extraction_metadata,
            )
            facts = extract_document_facts(
                files=prepared_files,
                tender_title=tender_title,
                tender_price_rub=tender_price_rub,
                policy=self.triage_policy,
            )
            extracted_signals = render_extracted_facts(facts)
            metadata_lines = [
                f"Название тендера: {tender_title.strip() or 'не указано'}",
                f"Цена тендера, руб.: {format_price_for_prompt(tender_price_rub)}",
            ]
            if facts.triage_signals:
                metadata_lines.append(f"Эвристика приоритета: {'; '.join(facts.triage_signals[:4])}")
            prompt = (
                f"Ссылка на тендер: {tender_url}\n\n"
                + "\n".join(metadata_lines)
                + "\n\n"
                f"{self.prompt_template}\n\n"
                "Извлечённые сигналы из документов. Используй их как вспомогательную выжимку, "
                "но не противоречь самим документам:\n"
                f"{extracted_signals}\n\n"
                "Дополнительное требование к формату ответа:\n"
                "Верни только JSON-объект без markdown-обёртки со следующими полями:\n"
                '{'
                '"decision": "Брать / Не брать / Уточнить", '
                '"confidence_percent": "целое число от 0 до 100", '
                '"summary_points": [], '
                '"llm_findings": [{"category": "risks|stack|requirements|docs|payment", "text": "краткий вывод", "quote": "точная цитата из документа", "source_file": "имя файла если известно"}], '
                '"analysis_markdown": "полный структурированный анализ по пунктам 1-5"'
                '}\n'
                f"Разрешённые значения decision: {', '.join(DECISIONS)}.\n"
                "summary_points не используй для итогового Excel: они будут построены системой из проверяемых фактов. "
                "В llm_findings добавляй только выводы, для которых можешь привести точную цитату из текста."
            )
            raw_text = self.provider.analyze_documents(prompt=prompt, files=prepared_files)
            budget_report = getattr(self.provider, "last_budget_report", None)
            if isinstance(budget_report, dict) and budget_report:
                extraction_report["llm_budget_report"] = budget_report
            parsed = self._parse_response(raw_text)
            parsed.llm_findings = verify_llm_findings(parsed.llm_findings, prepared_files)
            parsed.llm_raw_text = raw_text
            parsed.facts = facts
            parsed.extraction_report = extraction_report
            return postprocess_payload(parsed, facts, policy=self.triage_policy)
        except Exception as exc:
            error_type = "network_error" if _is_network_error_message(str(exc)) else "llm_error"
            if facts is not None:
                return _build_rule_based_fallback_payload(
                    facts=facts,
                    error_message=str(exc),
                    error_type=error_type,
                    extraction_report=locals().get("extraction_report", {}),
                )
            return AnalysisPayload(
                decision="Уточнить",
                confidence_percent=0,
                summary_points=_technical_retry_summary(),
                analysis_markdown=str(exc),
                completeness_label="низкая",
                llm_raw_text="",
                error_type=error_type,
            )
        finally:
            if temp_dir is not None:
                shutil.rmtree(temp_dir, ignore_errors=True)
                self._temp_dir = None

    def _prepare_files(self, files: list[Path]) -> list[Path]:
        prepared: list[Path] = []
        self._temp_dir = Path(tempfile.mkdtemp(prefix="tender_analysis_"))
        self._extraction_metadata = {}
        for path in files:
            suffix = path.suffix.lower()
            if suffix == ".docx":
                prepared.append(self._cached_or_convert(path, self._convert_docx_to_txt))
            elif suffix == ".doc":
                prepared.append(self._cached_or_convert(path, self._convert_doc_to_txt))
            elif suffix == ".xlsx":
                prepared.append(self._cached_or_convert(path, self._convert_xlsx_to_txt))
            elif suffix == ".xls":
                prepared.append(self._cached_or_convert(path, self._convert_xls_to_txt))
            elif suffix == ".pdf" and PdfReader is not None:
                prepared.append(self._cached_or_convert(path, self._convert_pdf_to_txt))
            else:
                self._record_extraction_metadata(path, path, converter="native_text", cache_hit=False)
                prepared.append(path)
        return prepared

    def _cached_or_convert(self, path: Path, converter) -> Path:
        cache_path = self._cache_path_for(path)
        if cache_path.exists() and _is_usable_cached_text(cache_path):
            prepared = self._temp_dir / f"{path.stem}.txt"
            shutil.copyfile(cache_path, prepared)
            self._record_extraction_metadata(
                path,
                prepared,
                converter=_converter_label(path),
                cache_hit=True,
                cached_from=str(cache_path),
            )
            return prepared
        if cache_path.exists():
            cache_path.unlink(missing_ok=True)
        converted = converter(path)
        try:
            text = converted.read_text(encoding="utf-8", errors="ignore")
            if not _is_usable_text(text):
                return converted
            cache_path.write_text(text, encoding="utf-8")
            metadata = dict(self._extraction_metadata.get(str(converted.resolve()), {}))
            metadata.update({"cache_hit": False, "cached_to": str(cache_path)})
            self._extraction_metadata[str(converted.resolve())] = metadata
            return converted
        except Exception:
            return converted

    def _cache_path_for(self, path: Path) -> Path:
        stat = path.stat()
        key = "|".join(
            [
                EXTRACTION_CACHE_VERSION,
                str(path.resolve()),
                str(int(stat.st_mtime)),
                str(stat.st_size),
                _file_sha256(path),
                str(self.max_chars_per_file),
            ]
        )
        digest = hashlib.sha1(key.encode("utf-8", errors="ignore")).hexdigest()
        return self.extract_cache_dir / f"{path.stem}__{digest}.txt"

    def _convert_docx_to_txt(self, path: Path) -> Path:
        out = self._temp_dir / f"{path.stem}.txt"
        with zipfile.ZipFile(path) as archive:
            xml_bytes = archive.read("word/document.xml")
        root = ElementTree.fromstring(xml_bytes)
        texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
        raw_text = "\n".join(filter(None, texts))
        out.write_text(self._truncate_text(raw_text), encoding="utf-8")
        self._record_extraction_metadata(
            path,
            out,
            converter="docx_xml",
            raw_text_chars=len(raw_text),
            cache_hit=False,
        )
        return out

    def _convert_doc_to_txt(self, path: Path) -> Path:
        out = self._temp_dir / f"{path.stem}.txt"
        text = extract_doc_with_textutil(path)
        if not text:
            out.write_text("", encoding="utf-8")
            self._record_extraction_metadata(path, out, converter="textutil", raw_text_chars=0, cache_hit=False)
            return out
        if not text.strip():
            out.write_text("", encoding="utf-8")
            self._record_extraction_metadata(path, out, converter="textutil", raw_text_chars=0, cache_hit=False)
            return out
        out.write_text(self._truncate_text(text), encoding="utf-8")
        self._record_extraction_metadata(path, out, converter="textutil", raw_text_chars=len(text), cache_hit=False)
        return out

    def _convert_xlsx_to_txt(self, path: Path) -> Path:
        out = self._temp_dir / f"{path.stem}.txt"
        workbook = openpyxl.load_workbook(path, data_only=True)
        parts: list[str] = []
        for sheet in _select_relevant_worksheets(workbook.worksheets):
            parts.append(f"[sheet] {sheet.title}")
            for line in _extract_meaningful_xlsx_lines(sheet, max_lines=120):
                parts.append(line)
        raw_text = "\n".join(parts)
        out.write_text(self._truncate_text(raw_text), encoding="utf-8")
        self._record_extraction_metadata(
            path,
            out,
            converter="openpyxl",
            raw_text_chars=len(raw_text),
            sheets_total=len(workbook.worksheets),
            sheets_read=min(len(workbook.worksheets), 8),
            cache_hit=False,
        )
        return out

    def _convert_xls_to_txt(self, path: Path) -> Path:
        out = self._temp_dir / f"{path.stem}.txt"
        if xlrd is None:
            out.write_text("", encoding="utf-8")
            self._record_extraction_metadata(path, out, converter="xlrd_missing", raw_text_chars=0, cache_hit=False)
            return out
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        parts: list[str] = []
        sheet_names = _select_relevant_xls_sheet_names(workbook.sheet_names())
        for sheet_name in sheet_names:
            sheet = workbook.sheet_by_name(sheet_name)
            parts.append(f"[sheet] {sheet.name}")
            for line in _extract_meaningful_xls_lines(sheet, max_lines=120):
                parts.append(line)
        raw_text = "\n".join(parts)
        out.write_text(self._truncate_text(raw_text), encoding="utf-8")
        self._record_extraction_metadata(
            path,
            out,
            converter="xlrd",
            raw_text_chars=len(raw_text),
            sheets_total=len(workbook.sheet_names()),
            sheets_read=len(sheet_names),
            cache_hit=False,
        )
        return out

    def _convert_pdf_to_txt(self, path: Path) -> Path:
        out = self._temp_dir / f"{path.stem}.txt"
        reader = PdfReader(str(path))
        parts: list[str] = []
        total_pages = len(reader.pages)
        max_pdf_chars = max(self.max_chars_per_file * 20, 160_000)
        total_chars = 0
        pages_read = 0
        for index, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            pages_read = index
            if page_text.strip():
                parts.append(f"\n[page {index}]\n{page_text}")
            total_chars += len(page_text)
            if total_chars >= max_pdf_chars:
                break
        text = "\n".join(parts).strip()
        ocr_attempted = len(text) < 40
        if ocr_attempted:
            fallback = _extract_pdf_with_pdftotext(path)
            if fallback.strip():
                text = fallback
        out.write_text(self._truncate_text(text), encoding="utf-8")
        self._record_extraction_metadata(
            path,
            out,
            converter="pypdf_pdftotext",
            raw_text_chars=len(text),
            pages_total=total_pages,
            pages_read=pages_read,
            ocr_attempted=ocr_attempted,
            cache_hit=False,
        )
        return out

    def _truncate_text(self, text: str) -> str:
        normalized = text.strip()
        if len(normalized) <= self.max_chars_per_file:
            return normalized
        excerpt = self._build_relevant_excerpt(normalized)
        if len(excerpt) <= self.max_chars_per_file:
            return excerpt.rstrip() + "\n\n[truncated]"
        return excerpt[: self.max_chars_per_file].rstrip() + "\n\n[truncated]"

    def _build_relevant_excerpt(self, text: str) -> str:
        marker_groups = [
            [
                "сведения о системе",
                "назначение системы",
                "цель и задачи модификации",
                "архитектур",
                "стек технолог",
                "технологическ",
                "используемые технологии",
                "технические требования",
                "требования к системе",
                "программно-техническ",
                "интеграц",
                "api",
                "rest",
                "субд",
                "база данных",
                "postgres",
                "oracle",
                "java",
                "javascript",
                "react",
                "python",
                "php",
                "1с",
                "модификац",
                "адаптац",
                "существующ",
            ],
            [
                "требования к участник",
                "требования к поставщик",
                "требования к исполнител",
                "требования к подрядчик",
                "квалификац",
                "опыт",
                "аналогичн",
                "референс",
                "оборот",
                "выручк",
                "финансов",
                "штат",
                "специалист",
                "персонал",
                "сро",
                "лиценз",
                "сертификат",
                "аккредит",
            ],
            [
                "оплат",
                "аванс",
                "предоплат",
                "постоплат",
                "обеспечени",
                "банковск",
                "гарант",
                "порядок расчет",
                "порядок расчёт",
                "договор",
                "акт приемк",
                "акт приёмк",
            ],
        ]
        lowered = text.casefold().replace("ё", "е")
        parts: list[str] = [text[: min(len(text), 2200)].strip()]
        group_budget = max(1200, (self.max_chars_per_file - sum(len(part) for part in parts)) // 3)
        seen_windows: list[tuple[int, int]] = [(0, min(len(text), 2200))]
        for markers in marker_groups:
            group_parts: list[str] = []
            for start, end in _marker_windows(lowered, markers, text_length=len(text)):
                if _overlaps_existing((start, end), seen_windows):
                    continue
                chunk = text[start:end].strip()
                if not chunk:
                    continue
                group_parts.append(chunk)
                seen_windows.append((start, end))
                if sum(len(part) for part in group_parts) >= group_budget:
                    break
            if group_parts:
                parts.append("\n\n[section]\n")
                parts.append("\n\n".join(group_parts))

        trimmed_parts: list[str] = []
        total = 0
        for part in parts:
            if not part:
                continue
            remaining = self.max_chars_per_file - total
            if remaining <= 0:
                break
            trimmed_parts.append(part[:remaining])
            total += len(trimmed_parts[-1])
        excerpt = "".join(trimmed_parts).strip()
        if excerpt:
            return excerpt
        return text[: self.max_chars_per_file]

    def _record_extraction_metadata(self, source: Path, prepared: Path, **metadata: object) -> None:
        payload = {
            "source_path": str(source),
            "source_name": source.name,
            "prepared_path": str(prepared),
            **metadata,
        }
        self._extraction_metadata[str(prepared.resolve())] = payload

    def _parse_response(self, text: str) -> AnalysisPayload:
        try:
            payload = json.loads(text)
            decision = str(payload.get("decision", "Уточнить")).strip()
            if decision not in DECISIONS:
                decision = "Уточнить"
            confidence = int(payload.get("confidence_percent", 0))
            confidence = max(0, min(100, confidence))
            summary_points = payload.get("summary_points") or []
            if not isinstance(summary_points, list):
                summary_points = []
            normalized_points = [
                _normalize_summary_point(str(point).strip(), index)
                for index, point in enumerate(summary_points, start=1)
                if str(point).strip()
            ]
            while len(normalized_points) < 5:
                normalized_points.append(
                    f"{len(normalized_points) + 1}. Недостаточно данных в документах для дополнительного вывода."
                )
            normalized_points = normalized_points[:5]
            analysis = str(payload.get("analysis_markdown", "")).strip() or text.strip()
            return AnalysisPayload(
                decision=decision,
                confidence_percent=confidence,
                summary_points=normalized_points,
                analysis_markdown=analysis,
                completeness_label="средняя",
                llm_raw_text=text,
                llm_findings=parse_llm_findings(payload.get("llm_findings")),
            )
        except Exception:
            return AnalysisPayload(
                decision="Уточнить",
                confidence_percent=0,
                summary_points=_technical_retry_summary(),
                analysis_markdown=text.strip(),
                completeness_label="низкая",
                llm_raw_text=text,
                error_type="llm_error",
            )


def _normalize_summary_point(value: str, index: int) -> str:
    trimmed = value.lstrip("-• ").strip()
    trimmed = _strip_leading_numbering(trimmed)
    if not trimmed:
        return f"{index}. Недостаточно данных в документах."
    return f"{index}. {trimmed}"


def _strip_leading_numbering(value: str) -> str:
    cleaned = value.strip()
    while True:
        updated = re.sub(r"^\s*\d+\s*[\.\)]\s*", "", cleaned)
        if updated == cleaned:
            return cleaned
        cleaned = updated.strip()


def _technical_retry_summary() -> list[str]:
    return [
        "Риски и спорные моменты: не оценены, требуется повторный прогон анализа.",
        "Требуемый стек: автоматически не определён из-за технической ошибки анализа.",
        "Требования к участнику: автоматически не разобраны, требуется повторный прогон.",
        "Кол-во требуемой документации для предоставления к участию: автоматически не определено.",
        "Условия оплаты и обеспечения: автоматически не разобраны, требуется повторный прогон.",
    ]


def _build_rule_based_fallback_payload(
    *,
    facts,
    error_message: str,
    error_type: str = "",
    extraction_report: dict[str, object] | None = None,
) -> AnalysisPayload:
    summary_points = build_canonical_summary_points(facts, "Уточнить")
    analysis_lines = [
        "LLM-анализ не завершился из-за технической ошибки.",
        "Ниже сохранена предварительная выжимка только из rule-based extraction.",
        "",
        render_extracted_facts(facts),
        "",
        f"Техническая ошибка: {error_message}",
    ]
    return AnalysisPayload(
        decision="Уточнить",
        confidence_percent=_fallback_confidence(facts),
        summary_points=summary_points,
        analysis_markdown="\n".join(line for line in analysis_lines if line is not None),
        completeness_label=facts.completeness_label,
        facts=facts,
        llm_raw_text="",
        error_type=error_type,
        extraction_report=extraction_report or {},
    )


def _is_network_error_message(message: str) -> bool:
    lowered = message.lower()
    markers = [
        "dns resolution failed",
        "could not resolve host",
        "name or service not known",
        "nodename nor servname provided",
        "timed out",
        "timeout",
        "connection reset",
        "temporarily unavailable",
        "non-zero exit status 6",
        "non-zero exit status 7",
        "non-zero exit status 28",
        "server disconnected",
    ]
    return any(marker in lowered for marker in markers)


def _fallback_risks(facts) -> str:
    risks: list[str] = []
    if facts.completeness_label == "низкая":
        risks.append("неполный комплект документов")
    for signal in facts.triage_signals[:2]:
        if "строительство/монтаж" in signal:
            risks.append("строительно-монтажный контур, не чистая разработка ПО")
        elif "enterprise/коробочный контур" in signal:
            risks.append("enterprise/коробочный контур")
        elif "тяжёлый ИБ/регуляторный контур" in signal:
            risks.append("тяжёлый ИБ/регуляторный контур")
        elif "низкий приоритет: тендер на сайт / веб-тематику" in signal:
            risks.append("низкий приоритет по бюджету и типу web-проекта")
        elif "непрофильный тип закупки" in signal:
            risks.append(signal.replace("непрофильный тип закупки: ", "непрофильный предмет: "))
        else:
            risks.append(signal)
    if not risks:
        risks.append("предварительный вывод, LLM не ответил")
    return "; ".join(risks)


def _fallback_stack(facts) -> str:
    if facts.stack:
        return ", ".join(facts.stack[:4])
    return "явный стек не указан"


def _fallback_party_requirements(facts) -> str:
    items = (
        facts.participant_requirements.turnover_texts[:1]
        + facts.participant_requirements.experience_texts[:1]
        + facts.participant_requirements.required_roles[:2]
        + facts.participant_requirements.licenses_or_certs[:2]
        + facts.turnover_requirements[:1]
        + facts.project_requirements[:1]
        + facts.team_requirements[:1]
        + facts.licenses[:1]
    )
    if items:
        normalized: list[str] = []
        for item in items[:4]:
            lowered = item.lower()
            if "референс" in lowered:
                normalized.append("референс-лист по реализованным объектам")
            elif "наличие персонала" in lowered:
                normalized.append("подтверждённый профильный персонал")
            elif "аттестован" in lowered:
                normalized.append("аттестованный персонал")
            else:
                normalized.append(item)
        return "; ".join(normalized[:4])
    return "явные требования не выделены"


def _fallback_docs(facts) -> str:
    if facts.document_roles:
        parts: list[str] = []
        role_labels = {
            "тз": "ТЗ",
            "извещение": "извещение",
            "требования_к_участнику": "требования к участнику",
            "договор": "условия договора",
            "нмцк": "ценовой документ",
        }
        for role in ("тз", "извещение", "требования_к_участнику", "договор", "нмцк"):
            if facts.document_roles.get(role):
                parts.append(role_labels[role])
        if parts:
            return f"{facts.completeness_label}; есть {', '.join(parts)}"
    return f"полнота {facts.completeness_label}"


def _fallback_payment(facts) -> str:
    if facts.payment:
        payment = _summarize_payment_signals(facts.payment)
        if payment:
            return payment
    return "условия оплаты не извлечены автоматически"


def _fallback_confidence(facts) -> int:
    if facts.completeness_label == "высокая":
        return 45
    if facts.completeness_label == "средняя":
        return 35
    return 25


def _summarize_payment_signals(items: list[str]) -> str:
    normalized = " ; ".join(item.strip().lower() for item in items if item)
    parts: list[str] = []
    if "по факту выполн" in normalized:
        parts.append("постоплата по факту")
    if "30 календарных д" in normalized:
        parts.append("оплата 30 календарных дней")
    elif "30 рабочих д" in normalized:
        parts.append("оплата 30 рабочих дней")
    if "авансирован" in normalized and "не предусмотр" in normalized:
        parts.append("без аванса")
    if "банковск" in normalized and "гарант" in normalized:
        parts.append("есть банковская гарантия")

    deduped: list[str] = []
    for part in parts:
        if part not in deduped:
            deduped.append(part)
    return "; ".join(deduped[:3])


def _extract_pdf_with_pdftotext(path: Path) -> str:
    import shutil
    pdftotext = shutil.which("pdftotext")
    if not pdftotext:
        return _extract_pdf_with_ocr(path)
    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = Path(tmpdir) / "out.txt"
            subprocess.run(
                [pdftotext, "-layout", str(path), str(out_path)],
                check=False,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            if out_path.exists():
                text = out_path.read_text(encoding="utf-8", errors="ignore")
                if text.strip():
                    return text
    except Exception:
        return _extract_pdf_with_ocr(path)
    return _extract_pdf_with_ocr(path)


def _extract_pdf_with_ocr(path: Path) -> str:
    import shutil
    pdftoppm = shutil.which("pdftoppm")
    tesseract = shutil.which("tesseract")
    if not pdftoppm or not tesseract:
        return ""
    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            prefix = Path(tmpdir) / "page"
            subprocess.run(
                [pdftoppm, "-r", "200", "-f", "1", "-l", "5", "-png", str(path), str(prefix)],
                check=False,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            texts: list[str] = []
            for img in sorted(Path(tmpdir).glob("page-*.png"))[:5]:
                outbase = img.with_suffix("")
                subprocess.run(
                    [tesseract, str(img), str(outbase), "-l", "rus+eng", "--dpi", "200"],
                    check=False,
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
                txt = outbase.with_suffix(".txt")
                if txt.exists():
                    texts.append(txt.read_text(encoding="utf-8", errors="ignore"))
            return "\n".join(texts)
    except Exception:
        return ""
    return ""


def _is_usable_cached_text(path: Path) -> bool:
    try:
        return _is_usable_text(path.read_text(encoding="utf-8", errors="ignore"))
    except Exception:
        return False


def _is_usable_text(text: str) -> bool:
    return sum(1 for char in text if not char.isspace()) >= 80


def _select_relevant_worksheets(worksheets: list[openpyxl.worksheet.worksheet.Worksheet]):
    scored = sorted(worksheets, key=_worksheet_priority_key)
    return scored[: min(len(scored), 8)]


def _worksheet_priority_key(sheet) -> tuple[int, int, str]:
    title = _normalize_sheet_name(sheet.title)
    if any(token in title for token in ["тз", "техничес", "извещ", "документац", "требован", "кп", "коммерч", "квалифик"]):
        band = 0
    elif any(token in title for token in ["договор", "контракт", "нмц", "обоснован", "смет"]):
        band = 1
    else:
        band = 2
    return (band, sheet.max_row * sheet.max_column, sheet.title.lower())


def _extract_meaningful_xlsx_lines(sheet, max_lines: int) -> list[str]:
    lines: list[str] = []
    seen: set[str] = set()
    dense_row_seen = False
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 200), values_only=True):
        values = [_normalize_cell_value(cell) for cell in row if _normalize_cell_value(cell)]
        if not values:
            continue
        if len(values) >= 3:
            dense_row_seen = True
        if not dense_row_seen and len(values) == 1 and len(values[0]) < 3:
            continue
        line = " | ".join(values[:8]).strip()
        if not line or line in seen:
            continue
        seen.add(line)
        lines.append(line)
        if len(lines) >= max_lines:
            break
    return lines


def _normalize_cell_value(value: object) -> str:
    if value in (None, ""):
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text


def _normalize_sheet_name(value: str) -> str:
    return value.casefold().replace("ё", "е").strip()


def _select_relevant_xls_sheet_names(sheet_names: list[str]) -> list[str]:
    scored = sorted(sheet_names, key=lambda name: _worksheet_priority_key_name(name))
    return scored[: min(len(scored), 8)]


def _worksheet_priority_key_name(name: str) -> tuple[int, str]:
    title = _normalize_sheet_name(name)
    if any(token in title for token in ["тз", "техничес", "извещ", "документац", "требован", "кп", "коммерч", "квалифик"]):
        band = 0
    elif any(token in title for token in ["договор", "контракт", "нмц", "обоснован", "смет"]):
        band = 1
    else:
        band = 2
    return (band, name.lower())


def _extract_meaningful_xls_lines(sheet, max_lines: int) -> list[str]:
    lines: list[str] = []
    seen: set[str] = set()
    dense_row_seen = False
    row_limit = min(sheet.nrows, 200)
    for row_index in range(row_limit):
        values = [
            _normalize_cell_value(sheet.cell_value(row_index, col_index))
            for col_index in range(min(sheet.ncols, 16))
        ]
        values = [value for value in values if value]
        if not values:
            continue
        if len(values) >= 3:
            dense_row_seen = True
        if not dense_row_seen and len(values) == 1 and len(values[0]) < 3:
            continue
        line = " | ".join(values[:8]).strip()
        if not line or line in seen:
            continue
        seen.add(line)
        lines.append(line)
        if len(lines) >= max_lines:
            break
    return lines


def _marker_windows(lowered_text: str, markers: list[str], *, text_length: int) -> list[tuple[int, int]]:
    windows: list[tuple[int, int]] = []
    for marker in markers:
        start = 0
        while True:
            idx = lowered_text.find(marker, start)
            if idx < 0:
                break
            windows.append((max(0, idx - 450), min(text_length, idx + 1800)))
            start = idx + len(marker)
    return _merge_windows(windows)


def _merge_windows(windows: list[tuple[int, int]]) -> list[tuple[int, int]]:
    merged: list[tuple[int, int]] = []
    for start, end in sorted(windows):
        if not merged or start > merged[-1][1] + 200:
            merged.append((start, end))
        else:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
    return merged


def _overlaps_existing(window: tuple[int, int], existing: list[tuple[int, int]]) -> bool:
    start, end = window
    return any(start < other_end and end > other_start for other_start, other_end in existing)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as fh:
            for chunk in iter(lambda: fh.read(1024 * 1024), b""):
                digest.update(chunk)
    except Exception:
        return ""
    return digest.hexdigest()


def _converter_label(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return "docx_xml"
    if suffix == ".doc":
        return "textutil"
    if suffix == ".xlsx":
        return "openpyxl"
    if suffix == ".xls":
        return "xlrd"
    if suffix == ".pdf":
        return "pypdf_pdftotext"
    return "native_text"
