from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment

from tender_agent.analysis_types import AnalysisPayload, ExtractedFacts, FactHit, LLMFinding, ParticipantRequirements
from tender_agent.analysis import TenderAnalyzer
from tender_agent.document_facts import count_hits
from tender_agent.document_facts import _extract_min_turnover_rub, extract_license_signals, extract_stack_signals
from tender_agent.evidence import collect_summary_evidence
from tender_agent.export.excel_migration import cleanup_legacy_workbook, remove_technical_rows, sort_sheet_by_deadline
from tender_agent.export.excel_schema import column_count, column_index, ensure_sheet_structure
from tender_agent.quality_gates import apply_quality_gates
from tender_agent.local_review import _analysis_role, _select_analysis_files
from tender_agent.llm_findings import merge_verified_findings_into_summary, verify_llm_findings
from tender_agent.triage_rules import postprocess_payload


class AuditRegressionTests(unittest.TestCase):
    def test_turnover_30_mln_is_parsed(self) -> None:
        self.assertEqual(_extract_min_turnover_rub(["оборот не менее 30 млн рублей"]), 30_000_000)

    def test_short_st_ld_do_not_trigger_without_industrial_context(self) -> None:
        text = "best effort rest status standard"
        self.assertNotIn("ST", extract_stack_signals(text))
        self.assertNotIn("LD", extract_stack_signals(text))

    def test_short_st_ld_trigger_with_industrial_context(self) -> None:
        text = "разработка scada по мэк 61131-3 на языках st и ld"
        stack = extract_stack_signals(text)
        self.assertIn("SCADA", stack)
        self.assertIn("ST", stack)
        self.assertIn("LD", stack)

    def test_license_agreement_is_not_participant_license(self) -> None:
        text = "исполнитель получает право использования по лицензионному соглашению на по заказчика"
        self.assertNotIn("лицензия", extract_license_signals(text))

    def test_otzyv_file_name_is_not_tz(self) -> None:
        from pathlib import Path

        self.assertNotEqual(_analysis_role(Path("Отзыв на запрос.docx")), "тз")

    def test_postprocess_preserves_error_and_extraction_report(self) -> None:
        facts = _minimal_facts()
        payload = AnalysisPayload(
            decision="Уточнить",
            confidence_percent=50,
            summary_points=["1. Риски: x"],
            analysis_markdown="x",
            completeness_label="низкая",
            facts=facts,
            error_type="llm_error",
            extraction_report={"files": [{"status": "ok"}]},
        )
        result = postprocess_payload(payload, facts)
        self.assertEqual(result.error_type, "llm_error")
        self.assertEqual(result.extraction_report, {"files": [{"status": "ok"}]})

    def test_evidence_uses_fact_anchor_without_fake_page(self) -> None:
        facts = _minimal_facts()
        facts.stack = ["PostgreSQL"]
        facts.fact_hits = {
            "stack": [
                FactHit(
                    field="stack",
                    term="PostgreSQL",
                    file_name="ТЗ.docx",
                    fragment="СУБД PostgreSQL используется как основная БД",
                    offset=42,
                )
            ]
        }
        evidence = collect_summary_evidence(facts=facts, files=[])
        rendered = evidence["stack"].render()
        self.assertIn("ТЗ.docx", rendered)
        self.assertNotIn("страница 1", rendered)

    def test_extraction_gate_caps_confidence_for_empty_key_file(self) -> None:
        facts = _minimal_facts()
        payload = AnalysisPayload(
            decision="Брать",
            confidence_percent=90,
            summary_points=["1. Стек: Python", "2. Требования к контрагенту: не указано"],
            analysis_markdown="",
            completeness_label="низкая",
            facts=facts,
            extraction_report={
                "files": [{"source_name": "Техническое задание.doc", "status": "empty_text"}],
            },
        )
        result = apply_quality_gates(payload, facts)
        self.assertEqual(result.decision, "Уточнить")
        self.assertLessEqual(result.confidence_percent, 40)
        self.assertIn("extraction_empty_key_file", facts.quality_flags)

    def test_budget_drop_caps_confidence(self) -> None:
        facts = _minimal_facts()
        payload = AnalysisPayload(
            decision="Брать",
            confidence_percent=90,
            summary_points=["1. Стек: Python", "2. Требования к контрагенту: не указано"],
            analysis_markdown="",
            completeness_label="средняя",
            facts=facts,
            extraction_report={
                "files": [{"source_name": "ТЗ.docx", "status": "ok"}],
                "llm_budget_report": {"dropped_files": [{"name": "Договор.docx"}]},
            },
        )
        result = apply_quality_gates(payload, facts)
        self.assertLessEqual(result.confidence_percent, 60)
        self.assertIn("files_dropped_by_llm_budget", facts.quality_flags)

    def test_partial_pdf_and_truncated_key_file_cap_confidence(self) -> None:
        facts = _minimal_facts()
        payload = AnalysisPayload(
            decision="Брать",
            confidence_percent=95,
            summary_points=["1. Стек: Python", "2. Требования к контрагенту: не указано"],
            analysis_markdown="",
            completeness_label="средняя",
            facts=facts,
            extraction_report={
                "files": [
                    {
                        "source_name": "Техническое задание.pdf",
                        "status": "ok",
                        "truncated": True,
                        "pages_total": 40,
                        "pages_read": 12,
                    }
                ],
            },
        )
        result = apply_quality_gates(payload, facts)
        self.assertLessEqual(result.confidence_percent, 60)
        self.assertIn("key_pdf_partially_read", facts.quality_flags)
        self.assertIn("key_file_truncated_before_llm", facts.quality_flags)

    def test_llm_finding_verification_requires_quote(self) -> None:
        from pathlib import Path
        import tempfile

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "doc.txt"
            path.write_text("Требуется опыт исследований в категории геосервисов.", encoding="utf-8")
            findings = [
                LLMFinding(
                    category="requirements",
                    text="опыт исследований в категории геосервисов",
                    quote="опыт исследований в категории геосервисов",
                ),
                LLMFinding(
                    category="stack",
                    text="React",
                    quote="React",
                ),
            ]
            verified = verify_llm_findings(findings, [path])
            self.assertTrue(verified[0].verified)
            self.assertFalse(verified[1].verified)

    def test_verified_llm_finding_can_replace_generic_summary_line(self) -> None:
        summary = [
            "1. Риски: существенные риски не выявлены",
            "2. Стек: явный стек не извлечён",
            "3. Требования к контрагенту: не указано",
        ]
        findings = [
            LLMFinding(
                category="requirements",
                text="опыт исследований в категории геосервисов",
                quote="опыт исследований в категории геосервисов",
                verified=True,
            )
        ]
        merged = merge_verified_findings_into_summary(summary, findings)
        self.assertIn("опыт исследований в категории геосервисов", merged[2])

    def test_evidence_renderer_keeps_verified_llm_summary(self) -> None:
        facts = _minimal_facts()
        facts.stack = []
        summary = [
            "1. Риски: существенные риски не выявлены",
            "2. Стек: Python; React",
            "3. Требования к контрагенту: опыт исследований в категории геосервисов",
        ]
        rendered = collect_summary_evidence(facts=facts, files=[])
        self.assertEqual(rendered, {})
        from tender_agent.evidence import render_summary_points_with_evidence

        lines = render_summary_points_with_evidence(
            facts=facts,
            decision="Уточнить",
            files=[],
            deadline_at=datetime(2026, 5, 1),
            summary_points=summary,
        )
        self.assertEqual(lines[1], "2. Стек: Python; React")
        self.assertEqual(lines[2], "3. Требования к контрагенту: опыт исследований в категории геосервисов")

    def test_unverified_llm_finding_does_not_replace_summary(self) -> None:
        summary = ["3. Требования к контрагенту: не указано"]
        findings = [
            LLMFinding(
                category="requirements",
                text="React",
                quote="React",
                verified=False,
            )
        ]
        self.assertEqual(merge_verified_findings_into_summary(summary, findings), summary)

    def test_short_policy_tokens_require_word_boundaries(self) -> None:
        self.assertEqual(count_hits("организация конференции", ["ии"]), 0)
        self.assertEqual(count_hits("подготовка презентации", ["pr"]), 0)
        self.assertEqual(count_hits("ai платформа", ["ai"]), 1)

    def test_selection_limit_one_still_keeps_role_coverage(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            files = [
                root / "Техническое задание.docx",
                root / "Требования к участникам.docx",
                root / "Проект договора.docx",
                root / "Обоснование НМЦ.xlsx",
                root / "Извещение.pdf",
            ]
            for path in files:
                path.write_text("placeholder", encoding="utf-8")
            selected = _select_analysis_files(files, limit=1)
            self.assertGreaterEqual(len(selected), 5)

    def test_xls_conversion_extracts_text(self) -> None:
        import xlwt

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            prompt = root / "prompt.md"
            prompt.write_text("stub", encoding="utf-8")
            xls_path = root / "Требования.xls"
            workbook = xlwt.Workbook()
            sheet = workbook.add_sheet("Требования")
            sheet.write(0, 0, "оборот не менее 30 млн рублей")
            sheet.write(1, 0, "опыт проведения исследований в категории геосервисов")
            workbook.save(str(xls_path))

            analyzer = TenderAnalyzer(
                provider_name="stub",
                api_key="",
                model="stub",
                base_url="",
                prompt_template_path=prompt,
                max_chars_per_file=8000,
            )
            prepared = analyzer._prepare_files([xls_path])
            text = prepared[0].read_text(encoding="utf-8")
            self.assertIn("оборот не менее 30 млн рублей", text)
            self.assertIn("опыт проведения исследований", text)

    def test_sort_preserves_user_columns_status_and_comment(self) -> None:
        wb = Workbook()
        ws = wb.active
        ensure_sheet_structure(ws)
        extra_col = column_count() + 1
        ws.cell(row=1, column=extra_col, value="Ручная колонка")
        ws.append([
            "",
            "Проверено",
            "summary late",
            "late",
            datetime(2026, 5, 2),
            "ручной комментарий",
            "batch",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "manual late",
        ])
        ws.cell(row=2, column=extra_col).comment = Comment("ручное примечание", "user")
        ws.append([
            "",
            "",
            "summary early",
            "early",
            datetime(2026, 5, 1),
            "",
            "batch",
        ])

        sort_sheet_by_deadline(ws)

        self.assertEqual(ws.cell(row=2, column=column_index("title")).value, "early")
        self.assertEqual(ws.cell(row=3, column=column_index("title")).value, "late")
        self.assertEqual(ws.cell(row=3, column=column_index("status")).value, "Проверено")
        self.assertEqual(ws.cell(row=3, column=column_index("comment")).value, "ручной комментарий")
        self.assertEqual(ws.cell(row=3, column=extra_col).value, "manual late")
        self.assertIsNotNone(ws.cell(row=3, column=extra_col).comment)
        self.assertEqual(ws.cell(row=3, column=extra_col).comment.text, "ручное примечание")

    def test_cleanup_does_not_rebuild_just_because_user_extra_column_exists(self) -> None:
        wb = Workbook()
        ws = wb.active
        ensure_sheet_structure(ws)
        extra_col = column_count() + 1
        ws.cell(row=1, column=extra_col, value="Ручная колонка")
        ws.append([
            "",
            "Проверено",
            "summary",
            "title",
            datetime(2026, 5, 1),
            "comment",
            "batch",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "manual value",
        ])

        cleanup_legacy_workbook(wb)

        self.assertEqual(ws.cell(row=1, column=extra_col).value, "Ручная колонка")
        self.assertEqual(ws.cell(row=2, column=extra_col).value, "manual value")

    def test_remove_technical_rows_keeps_rows_with_user_comment(self) -> None:
        wb = Workbook()
        ws = wb.active
        ensure_sheet_structure(ws)
        ws.append([
            "",
            "",
            "требуется повторный прогон анализа",
            "title",
            datetime(2026, 5, 1),
            "не удалять",
        ])

        remove_technical_rows(wb)

        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.cell(row=2, column=column_index("comment")).value, "не удалять")

def _minimal_facts() -> ExtractedFacts:
    return ExtractedFacts(
        procurement_type="разработка",
        key_files=[],
        stack=[],
        turnover_requirements=[],
        project_requirements=[],
        team_requirements=[],
        licenses=[],
        payment=[],
        document_roles={},
        completeness_label="низкая",
        completeness_notes=[],
        triage_signals=[],
        participant_requirements=ParticipantRequirements(),
    )


if __name__ == "__main__":
    unittest.main()
