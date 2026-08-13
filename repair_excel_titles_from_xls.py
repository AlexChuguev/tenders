from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from tender_agent.config import Settings
from tender_agent.excel_loader import load_tenders
from tender_agent.export.excel_schema import column_index


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: .venv/bin/python repair_excel_titles_from_xls.py <source_xls> <output_xlsx>")
        return 2

    source_xls = Path(sys.argv[1]).expanduser().resolve()
    output_xlsx = Path(sys.argv[2]).expanduser().resolve()

    base = Path("/Users/alexchuguev/Documents/tenders")
    settings = Settings.load(base)

    tenders = load_tenders(
        xls_path=source_xls,
        url_column=settings.platform_tender_url_column,
        id_column=settings.platform_tender_id_column,
        title_column=settings.platform_tender_title_column,
    )
    url_to_title = {t.url.strip(): t.title for t in tenders if t.url}
    title_to_url = {t.title.strip(): t.url.strip() for t in tenders if t.title and t.url}
    title_deadline_to_url = {
        (t.title.strip(), t.deadline_at): t.url.strip()
        for t in tenders
        if t.title and t.url and t.deadline_at
    }

    workbook = load_workbook(output_xlsx)
    changed = 0
    for worksheet in workbook.worksheets:
        for row_index in range(5, worksheet.max_row + 1):
            title_cell = worksheet.cell(row=row_index, column=column_index("title"))
            current_title = str(title_cell.value or "").strip()
            hyperlink = title_cell.hyperlink
            if hyperlink and hyperlink.target:
                url = hyperlink.target.strip()
                expected_title = url_to_title.get(url)
                if expected_title and current_title != expected_title:
                    title_cell.value = expected_title
                    current_title = expected_title
                    changed += 1
            deadline_cell = worksheet.cell(row=row_index, column=column_index("deadline")).value
            expected_url = title_deadline_to_url.get((current_title, deadline_cell)) or title_to_url.get(current_title)
            if expected_url:
                current_url = hyperlink.target.strip() if hyperlink and hyperlink.target else ""
                if current_url != expected_url:
                    title_cell.hyperlink = expected_url
                    title_cell.style = "Hyperlink"
                    changed += 1

    workbook.save(output_xlsx)
    print(f"changed={changed}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
