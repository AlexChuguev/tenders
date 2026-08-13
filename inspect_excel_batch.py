from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: .venv/bin/python inspect_excel_batch.py <excel_path> <batch_name>")
        return 2
    excel_path = Path(sys.argv[1]).expanduser().resolve()
    batch_name = sys.argv[2].strip()
    if not excel_path.exists():
        print(f"Excel not found: {excel_path}")
        return 2
    if not batch_name:
        print("Batch name is required.")
        return 2

    workbook = load_workbook(excel_path)
    missing_status = []
    missing_summary = []
    total = 0

    for worksheet in workbook.worksheets:
        headers = {
            str(worksheet.cell(row=1, column=col).value or "").strip(): col
            for col in range(1, worksheet.max_column + 1)
        }
        batch_col = headers.get("Выгрузка")
        summary_col = headers.get("Summary")
        status_col = headers.get("Статус")
        title_col = headers.get("Наименование лота")
        if not all([batch_col, summary_col, status_col, title_col]):
            continue
        for row_index in range(2, worksheet.max_row + 1):
            batch_value = str(worksheet.cell(row=row_index, column=batch_col).value or "").strip()
            if batch_value != batch_name:
                continue
            total += 1
            status_value = str(worksheet.cell(row=row_index, column=status_col).value or "").strip()
            summary_value = str(worksheet.cell(row=row_index, column=summary_col).value or "").strip()
            title_value = str(worksheet.cell(row=row_index, column=title_col).value or "").strip()
            if not status_value:
                missing_status.append((worksheet.title, row_index, title_value))
            if not summary_value:
                missing_summary.append((worksheet.title, row_index, title_value))

    print(f"Batch: {batch_name}")
    print(f"Total rows: {total}")
    print(f"Missing status: {len(missing_status)}")
    print(f"Missing summary: {len(missing_summary)}")
    if missing_status:
        print("Rows missing status:")
        for sheet, row, title in missing_status:
            print(f"  {sheet} #{row}: {title}")
    if missing_summary:
        print("Rows missing summary:")
        for sheet, row, title in missing_summary:
            print(f"  {sheet} #{row}: {title}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
